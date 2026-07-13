"""Bidirectional hyperhemispherical particle-filter diagnostics for AMASS poses.

The anomaly score for each joint-frame is the geodesic distance between the
observed quaternion and the pre-update HHPF prediction.  A frame is flagged
only when both the forward and backward filters consider it unlikely.
"""

from __future__ import annotations

import io
from pathlib import Path

import gdown
import matplotlib.pyplot as plt
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from pyrecest.distributions import (
    HyperhemisphericalDiracDistribution,
    HyperhemisphericalWatsonDistribution,
)
from pyrecest.filters import HyperhemisphericalParticleFilter


FILE_ID = "1772qebtt6dd3H4o0hyHROEnjBEB7BYH9"
PROJECT_ROOT = Path(__file__).resolve().parent.parent
OUTPUT_DIR = PROJECT_ROOT / "outputs"
RANDOM_SEED = 20260712
N_PARTICLES = 1000
PROCESS_NOISE_STD = 0.01
MEASUREMENT_KAPPA = 500.0
MAD_MULTIPLIER = 4.0
MIN_SCORE_DEG = 5.0
JOINT_IDX_FOR_PLOT = 16
EXCEL_REPORT_PATH = OUTPUT_DIR / "bidirectional_hhpf_incorrect_frames.xlsx"


def process_sequence(file_id: str):
    """Download and return an AMASS quaternion sequence."""
    buffer = io.BytesIO()
    gdown.download(id=file_id, output=buffer, quiet=False)
    buffer.seek(0)
    data = np.load(buffer, allow_pickle=True)
    print(f"Loaded keys: {data.files}")
    return data["poses_quat"]


def normalize_quat(q: np.ndarray) -> np.ndarray:
    """Normalize quaternions and choose the upper-hemisphere representative."""
    q = np.asarray(q, dtype=np.float64)
    norms = np.linalg.norm(q, axis=-1, keepdims=True)
    if np.any(norms == 0):
        raise ValueError("Quaternion input contains a zero-norm value.")
    q = q / norms
    q = np.array(q, copy=True)
    q *= np.where(q[..., -1:] < 0, -1.0, 1.0)
    return q


def amass_to_filter_quat(poses_amass: np.ndarray) -> np.ndarray:
    """Convert AMASS quaternion order ``(w, x, y, z)`` to ``(x, y, z, w)``."""
    return normalize_quat(poses_amass[..., [1, 2, 3, 0]])


def quat_multiply(q_left: np.ndarray, q_right: np.ndarray) -> np.ndarray:
    """Hamilton product for quaternions in ``(x, y, z, w)`` order."""
    x1, y1, z1, w1 = np.moveaxis(q_left, -1, 0)
    x2, y2, z2, w2 = np.moveaxis(q_right, -1, 0)
    return np.stack(
        (
            w1 * x2 + x1 * w2 + y1 * z2 - z1 * y2,
            w1 * y2 - x1 * z2 + y1 * w2 + z1 * x2,
            w1 * z2 + x1 * y2 - y1 * x2 + z1 * w2,
            w1 * w2 - x1 * x2 - y1 * y2 - z1 * z2,
        ),
        axis=-1,
    )


def quat_conjugate(q: np.ndarray) -> np.ndarray:
    q = np.array(q, copy=True)
    q[..., :3] *= -1
    return q


def quat_geodesic_distance(q1: np.ndarray, q2: np.ndarray) -> float:
    """Return the shortest angular distance between two orientations."""
    q1 = normalize_quat(q1)
    q2 = normalize_quat(q2)
    dot = float(np.clip(abs(np.dot(q1, q2)), 0.0, 1.0))
    return float(2.0 * np.arccos(dot))


def initialize_particles(q: np.ndarray) -> np.ndarray:
    particles = q + np.random.randn(N_PARTICLES, 4) * PROCESS_NOISE_STD
    return normalize_quat(particles)


def run_directional_hhpf(
    observations: np.ndarray,
    frame_order: np.ndarray,
    seed: int,
) -> tuple[np.ndarray, np.ndarray, np.ndarray]:
    """Run one causal HHPF pass over the supplied frame order.

    ``predicted`` is always saved before the observation update, so its
    residual measures an observation against a prediction that did not use
    that observation.  With reversed ``frame_order``, the same logic predicts
    from future original frames.
    """
    np.random.seed(seed)

    num_frames, num_joints, _ = observations.shape
    predicted = np.full((num_frames, num_joints, 4), np.nan)
    posterior = np.full((num_frames, num_joints, 4), np.nan)
    residuals = np.full((num_frames, num_joints), np.nan)
    particle_filters = [
        HyperhemisphericalParticleFilter(n_particles=N_PARTICLES, dim=3)
        for _ in range(num_joints)
    ]

    initial_frame = int(frame_order[0])
    for joint_idx, particle_filter in enumerate(particle_filters):
        q_initial = observations[initial_frame, joint_idx]
        particle_filter.set_state(
            HyperhemisphericalDiracDistribution(initialize_particles(q_initial))
        )
        posterior[initial_frame, joint_idx] = normalize_quat(
            particle_filter.filter_state.mean()
        )

    previous_posterior = None
    last_posterior = posterior[initial_frame].copy()

    for frame_idx in frame_order[1:]:
        frame_idx = int(frame_idx)
        current_posterior = np.zeros((num_joints, 4))

        if previous_posterior is None:
            increments = np.tile(np.array([0.0, 0.0, 0.0, 1.0]), (num_joints, 1))
        else:
            # Extrapolate the last filtered joint rotation.  This preserves
            # smooth, fast motion better than a pure identity random walk.
            increments = normalize_quat(
                quat_multiply(last_posterior, quat_conjugate(previous_posterior))
            )

        for joint_idx, particle_filter in enumerate(particle_filters):
            particles = quat_multiply(increments[joint_idx], particle_filter.filter_state.d)
            particles += np.random.randn(len(particles), 4) * PROCESS_NOISE_STD
            particle_filter.filter_state.d = normalize_quat(particles)

            q_prediction = normalize_quat(particle_filter.filter_state.mean())
            q_observed = observations[frame_idx, joint_idx]
            predicted[frame_idx, joint_idx] = q_prediction
            residuals[frame_idx, joint_idx] = quat_geodesic_distance(
                q_prediction, q_observed
            )

            # The score above is computed before this update.  Keeping the
            # measurement model fixed prevents the score from being altered
            # because the current observation was already labelled unusual.
            measurement = HyperhemisphericalWatsonDistribution(
                q_observed, kappa=MEASUREMENT_KAPPA
            )
            particle_filter.update_nonlinear_using_likelihood(measurement.pdf)
            current_posterior[joint_idx] = normalize_quat(
                particle_filter.filter_state.mean()
            )

        posterior[frame_idx] = current_posterior
        previous_posterior = last_posterior
        last_posterior = current_posterior

    return predicted, posterior, residuals


def robust_thresholds(scores: np.ndarray) -> np.ndarray:
    """Return one robust two-sided residual threshold per joint."""
    median = np.nanmedian(scores, axis=0)
    mad = np.nanmedian(np.abs(scores - median), axis=0)
    robust_sigma = 1.4826 * mad
    return np.maximum(
        median + MAD_MULTIPLIER * robust_sigma,
        np.deg2rad(MIN_SCORE_DEG),
    )


def run_bidirectional_hhpf(poses_amass: np.ndarray) -> dict[str, np.ndarray]:
    """Score each joint-frame from predictions made in both time directions."""
    observations = amass_to_filter_quat(poses_amass)
    num_frames = observations.shape[0]
    if num_frames < 3:
        raise ValueError("At least three frames are needed for bidirectional scoring.")

    forward_predicted, forward_posterior, forward_scores = run_directional_hhpf(
        observations, np.arange(num_frames), RANDOM_SEED
    )
    backward_predicted, backward_posterior, backward_scores = run_directional_hhpf(
        observations, np.arange(num_frames - 1, -1, -1), RANDOM_SEED + 1
    )

    # An isolated bad frame must be unlikely from both of its temporal sides.
    two_sided_scores = np.minimum(forward_scores, backward_scores)
    thresholds = robust_thresholds(two_sided_scores)
    anomaly_mask = two_sided_scores > thresholds[None, :]

    return {
        "observations": observations,
        "forward_predicted": forward_predicted,
        "backward_predicted": backward_predicted,
        "forward_posterior": forward_posterior,
        "backward_posterior": backward_posterior,
        "forward_scores": forward_scores,
        "backward_scores": backward_scores,
        "two_sided_scores": two_sided_scores,
        "thresholds": thresholds,
        "anomaly_mask": anomaly_mask,
    }


def print_flagged_frames(result: dict[str, np.ndarray]) -> None:
    mask = result["anomaly_mask"]
    thresholds = result["thresholds"]
    scores = result["two_sided_scores"]
    print("\nBidirectional HHPF candidates (observed vs predicted orientation):")
    print(f"Total candidate frames: {np.count_nonzero(mask.any(axis=1))}")
    for joint_idx in range(mask.shape[1]):
        frames = np.flatnonzero(mask[:, joint_idx])
        if len(frames) == 0:
            continue
        frame_scores_deg = np.degrees(scores[frames, joint_idx])
        print(
            f"  joint {joint_idx:2d}: {len(frames)} frames {frames.tolist()} "
            f"| threshold={np.degrees(thresholds[joint_idx]):.2f} deg "
            f"| scores={np.round(frame_scores_deg, 2).tolist()}"
        )


def style_header(worksheet, header_row: int, last_column: int) -> None:
    """Apply a compact, readable header style to a report sheet."""
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(bold=True, color="FFFFFF")
    for column_idx in range(1, last_column + 1):
        cell = worksheet.cell(header_row, column_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")


def export_incorrect_frames_excel(
    result: dict[str, np.ndarray],
    output_path: Path = EXCEL_REPORT_PATH,
) -> Path:
    """Write all joint-frame HHPF candidates to an auditable Excel workbook."""
    output_path.parent.mkdir(parents=True, exist_ok=True)

    mask = result["anomaly_mask"]
    num_frames, num_joints = mask.shape
    thresholds_deg = np.degrees(result["thresholds"])
    two_sided_scores_deg = np.degrees(result["two_sided_scores"])

    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "Summary"
    summary_sheet.sheet_view.showGridLines = False
    summary_sheet.append(["Metric", "Value"])
    summary_sheet.append(["Random seed", RANDOM_SEED])
    summary_sheet.append(["Frames processed", num_frames])
    summary_sheet.append(["Joints processed", num_joints])
    summary_sheet.append(["Candidate frames", int(np.count_nonzero(mask.any(axis=1)))])
    summary_sheet.append(["Candidate joint-frames", int(np.count_nonzero(mask))])
    summary_sheet.append(["Particles per joint", N_PARTICLES])
    summary_sheet.append(["Process noise standard deviation", PROCESS_NOISE_STD])
    summary_sheet.append(["Watson measurement kappa", MEASUREMENT_KAPPA])
    summary_sheet.append(["MAD multiplier", MAD_MULTIPLIER])
    summary_sheet.append(["Minimum score threshold (degrees)", MIN_SCORE_DEG])
    style_header(summary_sheet, header_row=1, last_column=2)
    summary_sheet.column_dimensions["A"].width = 36
    summary_sheet.column_dimensions["B"].width = 18
    summary_sheet.freeze_panes = "A2"

    joint_sheet = workbook.create_sheet("Joint Summary")
    joint_sheet.sheet_view.showGridLines = False
    joint_sheet.append(
        [
            "Joint index",
            "Candidate joint-frames",
            "Two-sided threshold (degrees)",
            "Maximum two-sided score (degrees)",
        ]
    )
    for joint_idx in range(num_joints):
        joint_scores = two_sided_scores_deg[:, joint_idx]
        finite_scores = joint_scores[np.isfinite(joint_scores)]
        max_score = float(np.max(finite_scores)) if len(finite_scores) else None
        joint_sheet.append(
            [
                joint_idx,
                int(np.count_nonzero(mask[:, joint_idx])),
                float(thresholds_deg[joint_idx]),
                max_score,
            ]
        )
    style_header(joint_sheet, header_row=1, last_column=4)
    joint_sheet.freeze_panes = "A2"
    joint_sheet.auto_filter.ref = f"A1:D{joint_sheet.max_row}"
    for column in ("C", "D"):
        joint_sheet.column_dimensions[column].width = 31
        for cell in joint_sheet[column][1:]:
            cell.number_format = "0.000"
    joint_sheet.column_dimensions["A"].width = 16
    joint_sheet.column_dimensions["B"].width = 25

    incorrect_sheet = workbook.create_sheet("Incorrect Frames")
    incorrect_sheet.sheet_view.showGridLines = False
    headers = [
        "Frame index",
        "Joint index",
        "Forward residual (degrees)",
        "Backward residual (degrees)",
        "Two-sided score (degrees)",
        "Joint threshold (degrees)",
        "Observed qx",
        "Observed qy",
        "Observed qz",
        "Observed qw",
        "Forward predicted qx",
        "Forward predicted qy",
        "Forward predicted qz",
        "Forward predicted qw",
        "Backward predicted qx",
        "Backward predicted qy",
        "Backward predicted qz",
        "Backward predicted qw",
    ]
    incorrect_sheet.append(headers)
    forward_scores_deg = np.degrees(result["forward_scores"])
    backward_scores_deg = np.degrees(result["backward_scores"])
    observations = result["observations"]
    forward_predicted = result["forward_predicted"]
    backward_predicted = result["backward_predicted"]

    for frame_idx, joint_idx in np.argwhere(mask):
        incorrect_sheet.append(
            [
                int(frame_idx),
                int(joint_idx),
                float(forward_scores_deg[frame_idx, joint_idx]),
                float(backward_scores_deg[frame_idx, joint_idx]),
                float(two_sided_scores_deg[frame_idx, joint_idx]),
                float(thresholds_deg[joint_idx]),
                *[float(value) for value in observations[frame_idx, joint_idx]],
                *[float(value) for value in forward_predicted[frame_idx, joint_idx]],
                *[float(value) for value in backward_predicted[frame_idx, joint_idx]],
            ]
        )

    style_header(incorrect_sheet, header_row=1, last_column=len(headers))
    incorrect_sheet.freeze_panes = "A2"
    incorrect_sheet.auto_filter.ref = f"A1:R{incorrect_sheet.max_row}"
    for column_idx in range(3, len(headers) + 1):
        for row_idx in range(2, incorrect_sheet.max_row + 1):
            incorrect_sheet.cell(row_idx, column_idx).number_format = "0.000000"
    for column_idx in range(1, len(headers) + 1):
        incorrect_sheet.column_dimensions[chr(64 + column_idx)].width = 23
    incorrect_sheet.column_dimensions["A"].width = 14
    incorrect_sheet.column_dimensions["B"].width = 14

    if incorrect_sheet.max_row > 1:
        table = Table(displayName="IncorrectFrames", ref=f"A1:R{incorrect_sheet.max_row}")
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False
        )
        incorrect_sheet.add_table(table)

    workbook.save(output_path)
    return output_path


def plot_joint_scores(result: dict[str, np.ndarray], joint_idx: int) -> None:
    """Save forward, backward, and two-sided prediction diagnostics."""
    num_frames = result["observations"].shape[0]
    frames = np.arange(num_frames)
    forward_deg = np.degrees(result["forward_scores"][:, joint_idx])
    backward_deg = np.degrees(result["backward_scores"][:, joint_idx])
    two_sided_deg = np.degrees(result["two_sided_scores"][:, joint_idx])
    threshold_deg = np.degrees(result["thresholds"][joint_idx])
    anomaly_mask = result["anomaly_mask"][:, joint_idx]

    fig, axes = plt.subplots(3, 1, figsize=(14, 10), sharex=True)
    panels = (
        (axes[0], forward_deg, "Forward pre-update prediction residual"),
        (axes[1], backward_deg, "Backward pre-update prediction residual"),
        (axes[2], two_sided_deg, "Two-sided score: min(forward, backward)"),
    )
    for axis, scores_deg, title in panels:
        axis.plot(frames, scores_deg, color="#378ADD", linewidth=1.2)
        axis.set_ylabel("Degrees")
        axis.set_title(title, fontsize=11)
        axis.set_ylim(bottom=0)
        axis.grid(True, linewidth=0.4, alpha=0.5, linestyle="--")
        axis.spines[["top", "right"]].set_visible(False)

    axes[2].axhline(
        threshold_deg,
        color="#E24B4A",
        linewidth=1.4,
        linestyle="--",
        label=f"robust threshold ({threshold_deg:.2f} deg)",
    )
    axes[2].scatter(
        frames[anomaly_mask],
        two_sided_deg[anomaly_mask],
        color="#E24B4A",
        s=45,
        zorder=3,
        label=f"candidate frames ({anomaly_mask.sum()})",
    )
    axes[2].legend(loc="upper right")
    axes[2].set_xlabel("Frame index")
    fig.suptitle(f"Bidirectional HHPF orientation diagnostics - joint {joint_idx}")
    fig.tight_layout()

    OUTPUT_DIR.mkdir(exist_ok=True)
    output_path = OUTPUT_DIR / f"bidirectional_hhpf_joint{joint_idx}.png"
    fig.savefig(output_path, dpi=150, bbox_inches="tight")
    print(f"Saved diagnostic plot: {output_path}")
    plt.show()


def main() -> None:
    poses_amass = process_sequence(FILE_ID)
    result = run_bidirectional_hhpf(poses_amass)
    report_path = export_incorrect_frames_excel(result)
    print(f"Saved Excel report: {report_path}")
    print_flagged_frames(result)
    plot_joint_scores(result, JOINT_IDX_FOR_PLOT)


if __name__ == "__main__":
    main()
